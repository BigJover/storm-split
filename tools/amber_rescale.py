"""One-shot Big Amber rescale (PLAN 10a T134, DECISIONS #385/#386/#394).

Multiplies every Amber cell of Storm-Split-Balance.xlsx by SCALE (1000) exactly once, and
appends the Tuning row `Amber scale` that records it. Run it twice and it refuses: the
`Amber scale` row is the marker. Every registered cell must hold its old value (a whole
number below SCALE) or the script stops before writing anything.

The registry below is every Amber column/row in the workbook. The exporter's own check
(export_constants.AMBER_FIELDS: every Amber value a multiple of the scale, none below it)
catches a column this registry ever misses.

Mastery's Amber cost is a formula (ROUND(base x growth^(L-1))): it is rewritten to round in
old units and then scale, so each level is exactly 1000x its old cost (the cumulative and
"matches to earn" columns follow).

    python3 tools/amber_rescale.py [workbook]
"""
import os
import sys

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(os.path.dirname(HERE), "Storm-Split-Balance.xlsx")

SCALE = 1000
SCALE_ROW_NAME = "Amber scale"
SCALE_NOTE = ("x every Amber value was multiplied by once (PLAN 10a T134, #394; tools/amber_rescale.py); "
              "saves older than this are migrated once by Shared/Profile (amberScale)")

# Tuning rows by name -> the old value they must hold.
TUNING_ROWS = {
    "Multiplayer loss payout": 5,
    "Typical amber per match": 50,
    "Mastery cost base": 25,
    "Competitive buy-in": 15,
    "Team win payout": 50,
    "Team loss payout": 5,
    "Royale 1st payout": 60,
    "Royale 2nd payout": 30,
    "Royale 3rd payout": 10,
    "Royale rest payout": 5,
}
# Tuning notes that quote an Amber number.
TUNING_NOTES = {
    "Competitive buy-in": ("Amber each player stakes (exporter: 10-20)",
                           "Amber each player stakes (exporter: 10,000-20,000)"),
}
# (sheet, column, header in row 4): one Amber value per named row from row 5 down (blank
# rows and Towers' "Cost per DPS" line skipped). Blank cells (note rows) are skipped.
COLUMNS = [
    ("Towers", 14, "Unlock cost (amber)"),
    ("Heroes", 24, "Unlock cost (amber)"),
    ("Difficulty", 9, "Clear reward (amber)"),
    ("Daily Haul", 2, "Amber"),
    ("Bounties", 7, "Amber"),
]
MASTERY_COST_COLUMN = 2  # B: =ROUND('Tuning'!$B$27*POWER('Tuning'!$B$28,An-1),0)
MASTERY_OLD = "=ROUND('Tuning'!$B$27*POWER('Tuning'!$B$28,A{r}-1),0)"
MASTERY_NEW = "=ROUND('Tuning'!$B$27/'Tuning'!$B${s}*POWER('Tuning'!$B$28,A{r}-1),0)*'Tuning'!$B${s}"


def fail(msg):
    sys.exit(f"amber_rescale: {msg} (nothing written)")


def old_amount(where, value):
    if not isinstance(value, (int, float)) or value != int(value) or value < 0 or value >= SCALE:
        fail(f"{where} holds {value!r}, not an old whole Amber value below {SCALE}")
    return int(value)


def tuning_rows(ws):
    rows = {}
    last = 0
    for r in range(1, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name is not None:
            rows[str(name).strip()] = r
            last = r
    return rows, last


def plan(wb):
    """Every (cell, old, new) change, checked before anything is written."""
    changes = []
    ws = wb["Tuning"]
    rows, last = tuning_rows(ws)
    if SCALE_ROW_NAME in rows:
        fail(f"Tuning already has '{SCALE_ROW_NAME}' (row {rows[SCALE_ROW_NAME]}): the sheet is rescaled; never run twice")
    for name, want in TUNING_ROWS.items():
        if name not in rows:
            fail(f"Tuning has no '{name}' row")
        cell = ws.cell(rows[name], 2)
        if cell.value != want:
            fail(f"Tuning '{name}' is {cell.value!r}, expected the old {want}")
        changes.append((cell, want, want * SCALE))
    for name, (old, new) in TUNING_NOTES.items():
        cell = ws.cell(rows[name], 3)
        if cell.value != old:
            fail(f"Tuning '{name}' note is {cell.value!r}, expected {old!r}")
        changes.append((cell, old, new))
    scale_row = last + 1
    for sheet, col, header in COLUMNS:
        s = wb[sheet]
        if str(s.cell(4, col).value).strip() != header:
            fail(f"{sheet} row 4 column {col} is {s.cell(4, col).value!r}, expected '{header}'")
        for r in range(5, s.max_row + 1):
            name = s.cell(r, 1).value
            if name is None or str(name).startswith("Cost per DPS"):  # Towers' gap rows / value.py line
                continue
            cell = s.cell(r, col)
            if cell.value is not None:
                old = old_amount(f"{sheet}!{cell.coordinate}", cell.value)
                changes.append((cell, old, old * SCALE))
    m = wb["Mastery"]
    if str(m.cell(4, MASTERY_COST_COLUMN).value).strip() != "Amber cost":
        fail("Mastery B4 is not 'Amber cost'")
    r = 5
    while m.cell(r, 1).value is not None and isinstance(m.cell(r, 1).value, (int, float)):
        cell = m.cell(r, MASTERY_COST_COLUMN)
        text = cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value
        if text != MASTERY_OLD.format(r=r):
            fail(f"Mastery!{cell.coordinate} is {text!r}, expected {MASTERY_OLD.format(r=r)!r}")
        new = MASTERY_NEW.format(r=r, s=scale_row)
        changes.append((cell, cell.value, ArrayFormula(cell.coordinate, new) if isinstance(cell.value, ArrayFormula) else new))
        r += 1
    return changes, scale_row


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(path)
    changes, scale_row = plan(wb)
    for cell, _old, new in changes:
        cell.value = new
    ws = wb["Tuning"]
    for col in range(1, 4):
        if ws.cell(scale_row, col).value is not None:
            fail(f"Tuning row {scale_row} is not empty")
    ws.cell(scale_row, 1).value = SCALE_ROW_NAME
    ws.cell(scale_row, 2).value = SCALE
    ws.cell(scale_row, 3).value = SCALE_NOTE
    wb.save(path)
    print(f"amber_rescale: x{SCALE} on {len(changes)} cells; Tuning row {scale_row} '{SCALE_ROW_NAME}' = {SCALE}")


if __name__ == "__main__":
    main()
