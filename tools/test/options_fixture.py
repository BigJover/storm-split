#!/usr/bin/env python3
"""
PLAN T113: prove every A/B/C option bar can fail. Builds a doctored COPY of the balance sheet
in a temp folder (the real sheet and Config.luau are never touched), appends B/C option rows
that should break each bar, and checks that tools/value.py --options exits 1 naming each one
and that the exporter's validation refuses each one. Run by tools/test.sh.

  dead option        Hunting Blind Trigger Happy T4 B: no step over tier 3
  must-pick          Hunting Blind Trigger Happy T4 C: Damage x 4
  Raider floor       Armory Plating T4 B, PvP effect: resist 26 (A: 35 + thorns) — passes the
                     slot bar (2 options: median = their mean) but not 0.8x the other option
  one Raider/tower   Field Hospital Ward T4 B and Tonic T4 B, both PvP effect (copies of A)
  worth caps         Falcon Roost (can't be damaged) Talons T5 B: Damage x 4 (too strong)
  hero option bar    Tracker Gunslinger T4 B: Fire rate x 3 (two options: A reads dead)
  hero parity        the same row: its T6 builds run far past 1.15x the A builds
"""

import contextlib
import io
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
ROOT = os.path.dirname(TOOLS)
sys.path.insert(0, TOOLS)
import export_constants  # noqa: E402

EXPECT = [
    "OPTIONS: Hunting Blind Trigger Happy T4 B Fixture Dead",
    "a dead option",
    "OPTIONS: Hunting Blind Trigger Happy T4 C Fixture Must",
    "a must-pick",
    "OPTIONS: Armory Plating T4 B Fixture Raider: Raider option",
    "Field Hospital: 2 Raider options",
    "WORTH: Falcon Roost Talons build",
    "too strong for a tower that cannot be damaged",
    "OPTIONS: Tracker Gunslinger T4 ",
    "PARITY: Tracker Gunslinger T6 builds",
]


def headers(ws):
    return {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}


def find(ws, entity, path, tier, option="A"):
    h = headers(ws)
    for r in range(5, ws.max_row + 1):
        if (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, h["Option"]).value) == (entity, path, tier, option):
            return r
    raise LookupError(f"{entity} {path} T{tier}{option}")


def append(ws, source, option, name, name_col, blank=False, raider=False, **steps):
    """A copy of row `source` at the bottom as option `option`, with `steps` (header -> value)."""
    h = headers(ws)
    r = ws.max_row + 1
    while ws.cell(r - 1, 1).value in (None, ""):
        r -= 1
    first_effect = h["Damage x"] if "Damage x" in h else h["Fire rate x"]
    for c in range(1, ws.max_column + 1):
        value = ws.cell(source, c).value
        if isinstance(value, ArrayFormula):
            value = value.text
        if blank and first_effect <= c < h["Option"] and c != name_col and not (isinstance(value, str) and value.startswith("=")):
            value = None
        ws.cell(r, c).value = value
    ws.cell(r, h["Option"]).value = option
    ws.cell(r, h["Blurb"]).value = "Fixture"
    ws.cell(r, name_col).value = name
    if raider:
        ws.cell(r, h["PvP effect"]).value = "yes"
    for head, value in steps.items():
        ws.cell(r, h[head]).value = value


def doctor(path):
    wb = openpyxl.load_workbook(path)
    tu, hu = wb["Tower Upgrades"], wb["Hero Upgrades"]
    name = headers(tu)["Name"]
    a = find(tu, "Hunting Blind", "Trigger Happy", 4)
    append(tu, a, "B", "Fixture Dead", name, blank=True)
    append(tu, a, "C", "Fixture Must", name, **{"Damage x": 4})
    append(tu, find(tu, "Armory", "Plating", 4), "B", "Fixture Raider", name, blank=True, raider=True, **{"Resist %": 26})
    append(tu, find(tu, "Field Hospital", "Ward", 4), "B", "Fixture Ward", name, raider=True)
    append(tu, find(tu, "Field Hospital", "Tonic", 4), "B", "Fixture Tonic", name, raider=True)
    append(tu, find(tu, "Falcon Roost", "Talons", 5), "B", "Fixture Eagle", name, **{"Damage x": 4})
    append(hu, find(hu, "Tracker", "Gunslinger", 4), "B", "Fixture Hero", headers(hu)["Name"], **{"Fire rate x": 3})
    wb.save(path)


def main():
    folder = tempfile.mkdtemp(prefix="options-fixture-")
    try:
        copy = os.path.join(folder, "doctored.xlsx")
        shutil.copy(export_constants.DEFAULT_XLSX, copy)
        doctor(copy)
        report = subprocess.run([sys.executable, os.path.join(TOOLS, "value.py"), "--options", copy],
                                capture_output=True, text=True)
        failed = []
        if report.returncode != 1:
            failed.append(f"value.py --options exited {report.returncode}, not 1:\n{report.stderr[-800:]}")
        for want in EXPECT:
            if want not in report.stdout:
                failed.append(f"value.py --options never says '{want}'")
        with contextlib.redirect_stdout(io.StringIO()):
            problems, _ = export_constants.validate(export_constants.read_data(copy))
        joined = "\n".join(problems)
        for want in EXPECT:
            if want not in joined:
                failed.append(f"the exporter doesn't refuse '{want}'")
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    if failed:
        print("options fixture: FAIL")
        for f in failed:
            print(f"  ! {f}")
        sys.exit(1)
    print(f"options fixture: every bar fails on the doctored copy ({len(EXPECT)} checks, value.py and the exporter)")


if __name__ == "__main__":
    main()
