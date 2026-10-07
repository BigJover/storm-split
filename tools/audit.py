#!/usr/bin/env python3
"""
Dino Hunters — Config <-> design-doc audit.

Catches drift between the approved docs, the spreadsheet and the code. Reads the
spreadsheet through the exporter (export_constants.read_data), never by parsing
Config.luau, and prints a findings list:

  1. names    every tier name in the UPGRADES.md / HEROES.md path tables equals the
              spreadsheet Name, keyed by (tower or hero, path, tier) and scoped to
              Towers vs Heroes (tier names repeat across the two)
  2. unlocks  unlock costs equal the DIRECTION.md ladder, extended by DECISIONS
              #13/#15/#16 once those rows exist
  3. dead     every Config field that is non-zero (true, non-empty) somewhere is
              referenced by name in some src/**/*.luau other than Config.luau, so no
              mechanic is silently unbuilt. Exempt, and listed as notes: Tuning levers
              that only sheet formulas read ("sheet-only", DECISIONS #28), and the
              DEAD_ALLOWED entries below
  4. species  every species in VISION.md's table exists with that display name
  5. perks    Mastery Perks (PLAN round 3 T32): every number in a perk's Text is one
              of that row's effect cells (Extra mark power x as a percentage), no Text
              says "pop" (DECISIONS #93), no effect column is blank on every row, no
              perk name equals a tower or hero tier name or a hero's ability name,
              whatever the case (DECISIONS #104), and each perk name is in HEROES.md
              "Mastery" (a note until the docs task)
  6. mastery  no Mastery sheet column header and no Config.Mastery field touches damage,
              fire rate, reload, recoil, spread, max HP or move speed (PLAN T51,
              DECISIONS #135: mastery perks are PvP-safe)

Usage:
    python3 tools/audit.py            report; exits 0 whatever it finds
    python3 tools/audit.py --strict   exits 1 if there are findings

tools/test.sh runs it after the specs.
"""

import contextlib
import glob
import io
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import export_constants  # noqa: E402

# Rows the DIRECTION.md ladder doesn't list yet, from DECISIONS.md. Checked only
# once the key exists in Config. (kind, key): (display, cost, source)
LADDER_EXTENSIONS = {
    ("hero", "MEDIC"): ("Field Medic", 75, "DECISIONS #13"),
    ("tower", "HOSPITAL"): ("Field Hospital", 100, "DECISIONS #15"),
    ("tower", "ARMORY"): ("Armory", 150, "DECISIONS #16"),
    # The round-4 batch: only the Amber price is 2x (DECISIONS #147, TOWERS_NEXT.md picks).
    ("tower", "COIL"): ("Storm Coil", 300, "DECISIONS #147"),
    ("tower", "FALCON"): ("Falcon Roost", 200, "DECISIONS #147"),
    ("tower", "TARPIT"): ("Tar Pit", 250, "DECISIONS #147"),
    ("tower", "BALLISTA"): ("Harpoon Ballista", 300, "DECISIONS #147"),
}


# Config fields no code reads, on purpose (DECISIONS #28). One reason each; nothing
# else belongs here. (block, field): reason
DEAD_ALLOWED = {
    ("Tuning", "StartingLives"): "superseded by Difficulty -> Starting lives",
    ("Enemies", "cashValue"): "display column; the code pays maxHp x CashPerEffectiveHP, the same value unrounded",
    ("Bounties", "title"): "display text: the Hunt Board card and the in-match toast (PLAN round 2 T28/T29)",
    ("Bounties", "text"): "display text: the Hunt Board card (PLAN round 2 T29)",
    ("Tuning", "MasteryPerkWorthCap"): "an exporter rule, not a game number: every perk's worth is checked against it (DECISIONS #99)",
}


# Config fields whose reader is a later task of the plan in progress. Listed as notes;
# an entry becomes a finding the moment some src file reads the field, so it can't
# outlive its task. (block, field): the task that reads it
DEAD_PENDING = {
    # Phase 7 batch 1 (PLAN T73): the levers whose readers come later in phase 7.
    ("Tuning", "PvPTowerDamage"): "T78 (PvP damage)",
    ("Tuning", "PvPHunterDamage"): "T78 (PvP damage)",
    ("Tuning", "SpawnShield"): "T78 (PvP damage)",
    ("Tuning", "CompetitiveBuyIn"): "T80 (Progression passes it to Stakes.canStake and escrows it)",
    ("Tuning", "CompetitiveEdgeMax"): "T84 (value.py/exporter edge bar; then DEAD_ALLOWED as an exporter rule)",
    ("Tuning", "PerkParityBand"): "T84 (value.py/exporter parity bar; then DEAD_ALLOWED as an exporter rule)",
}

# Words no tower or tier name may use: "Trophy" is reserved for ranked (DECISIONS #130,
# #139); "split" and "pop" were retired with balloons (#93, TOWERS_NEXT "Final names").
BANNED_NAME_WORDS = ("trophy", "split", "pop")


def read(name):
    with open(os.path.join(ROOT, name), encoding="utf-8") as f:
        return f.read()


def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def plain(text):
    """A table cell without markdown emphasis."""
    return text.replace("**", "").replace("*", "").strip()


# ---------------------------------------------------------------- 1. names

def doc_path_tables(text, heading_level):
    """{section display name: [(path id, [tier names]), ...]} from a design doc.

    A section starts at a heading of `heading_level` ('##' or '###') whose title is
    '<Display> (<note>) — ...'; its path rows are table rows whose first cell is bold.
    """
    sections = {}
    current = None
    heading = re.compile(r"^" + re.escape(heading_level) + r" (.+?) \(")
    for line in text.splitlines():
        if line.startswith("#"):
            m = heading.match(line)
            current = m.group(1).strip() if m else None
            if current:
                sections[current] = []
        elif current and line.startswith("|"):
            row = cells(line)
            if len(row) >= 3 and row[0].startswith("**"):
                sections[current].append((plain(row[0]), [plain(c) for c in row[2:]]))
    return {k: v for k, v in sections.items() if v}


def check_names(findings, block, items, doc_name, heading_level, compared):
    sections = doc_path_tables(read(doc_name), heading_level)
    by_display = {item["display"]: key for key, item in items.items()}
    for display, rows in sections.items():
        key = by_display.get(display)
        if key is None:
            findings.append(f"{doc_name} section '{display}' has no {block} with that display name")
            continue
        paths = items[key]["paths"]
        if len(rows) != len(paths):
            findings.append(f"{block}.{key}: {doc_name} lists {len(rows)} paths, the sheet has {len(paths)}")
        for p, (doc_id, doc_tiers) in enumerate(rows, start=1):
            if p > len(paths):
                break
            path = paths[p - 1]
            where = f"{block}.{key} path {p}"
            if doc_id != path["id"]:
                findings.append(f"{where}: {doc_name} path '{doc_id}', sheet '{path['id']}'")
            tiers = path["tiers"]
            if len(doc_tiers) != len(tiers):
                findings.append(f"{where} ({path['id']}): {doc_name} has {len(doc_tiers)} tiers, the sheet {len(tiers)}")
            for t, (doc_tier, tier) in enumerate(zip(doc_tiers, tiers), start=1):
                compared.append(1)
                if doc_tier != tier.get("name"):
                    findings.append(f"{where} ({path['id']}) T{t}: {doc_name} '{doc_tier}', sheet '{tier.get('name')}'")
    documented = set(sections)
    for key, item in items.items():
        if item["display"] not in documented:
            findings.append(f"{block}.{key} ({item['display']}) has no path table in {doc_name}")
    # Tower tier names: unique across every tower, none with a banned word (PLAN T62/T63).
    if block == "Towers":
        seen = {}
        for key, item in items.items():
            names = [(item["display"], "name")] + [(tier.get("name") or "", f"{path['id']} T{t}")
                                                   for path in item["paths"] for t, tier in enumerate(path["tiers"], start=1)]
            for name, where in names:
                words = re.findall(r"[a-z]+", name.lower())
                for bad in BANNED_NAME_WORDS:
                    if bad in words:
                        findings.append(f"{block}.{key} {where} '{name}' uses the banned word '{bad}'")
                if where != "name":
                    if name.lower() in seen:
                        findings.append(f"{block}.{key} {where} '{name}' repeats {seen[name.lower()]}")
                    seen.setdefault(name.lower(), f"{key} {where}")


# ---------------------------------------------------------------- 2. unlocks

def direction_ladder():
    """[(display, cost)] from the DIRECTION.md '- Unlocks:' bullet."""
    lines = read("DIRECTION.md").splitlines()
    for i, line in enumerate(lines):
        if line.startswith("- Unlocks:"):
            text = line[len("- Unlocks:"):]
            for more in lines[i + 1:]:
                if not more.startswith("  "):
                    break
                text += " " + more.strip()
            break
    else:
        return None
    # The ladder is the bullet's first sentence; later sentences are notes
    # ("The next tower batch costs about 2× (2026-10-02).").
    text = re.split(r"\.\s", text.strip(), maxsplit=1)[0]
    ladder = []
    for part in text.strip().rstrip(".").split(";"):
        part = part.strip()
        if part.endswith(" free"):
            for name in re.split(r"\s*[+/]\s*", part[: -len(" free")]):
                ladder.append((name.strip(), 0))
        else:
            for item in part.split(","):
                m = re.match(r"^(.+?)\s+(\d+)$", item.strip())
                if m:
                    ladder.append((m.group(1).strip(), int(m.group(2))))
                else:
                    ladder.append((item.strip(), None))
    return ladder


def check_unlocks(findings, data):
    ladder = direction_ladder()
    if ladder is None:
        findings.append("DIRECTION.md has no '- Unlocks:' bullet")
        return
    kinds = {"tower": ("Towers", data["Towers"]), "hero": ("Heroes", data["Heroes"])}
    expected = {}  # (kind, key) -> (cost, source)
    for display, cost in ladder:
        matches = [(kind, key) for kind, (_, items) in kinds.items()
                   for key, item in items.items() if item["display"] == display]
        if cost is None:
            findings.append(f"DIRECTION.md ladder entry '{display}' has no cost")
        elif not matches:
            findings.append(f"DIRECTION.md ladder names '{display}', which is no tower or hero in the sheet")
        elif len(matches) > 1:
            findings.append(f"DIRECTION.md ladder name '{display}' is both a tower and a hero")
        else:
            expected[matches[0]] = (cost, "DIRECTION.md")
    for (kind, key), (display, cost, source) in LADDER_EXTENSIONS.items():
        items = kinds[kind][1]
        if key in items:
            expected[(kind, key)] = (cost, source)
            if items[key]["display"] != display:
                findings.append(f"{kinds[kind][0]}.{key} is '{items[key]['display']}', {source} says '{display}'")
    for kind, (block, items) in kinds.items():
        for key, item in items.items():
            want = expected.get((kind, key))
            if want is None:
                findings.append(f"{block}.{key} ({item['display']}) is on no unlock ladder")
            elif item.get("unlockCost") != want[0]:
                findings.append(f"{block}.{key} ({item['display']}) unlock {item.get('unlockCost')}, {want[1]} says {want[0]}")


# ---------------------------------------------------------------- 3. dead columns

def live(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    if isinstance(value, str):
        return value != ""
    return value is not None


def config_fields(data):
    """{(block, field): [where it's non-zero]} for every leaf field in Config."""
    fields = {}

    def note(block, row, where):
        for field, value in row.items():
            if isinstance(value, (dict, list)):
                continue
            entry = fields.setdefault((block, field), [])
            if live(value):
                entry.append(where)

    note("Tuning", data["Tuning"], "Tuning")
    for block in ("Towers", "Heroes"):
        for key, item in data[block].items():
            note(block, item, key)
            for p, path in enumerate(item["paths"], start=1):
                note(f"{block} paths", path, f"{key} path {p}")
                for t, tier in enumerate(path["tiers"], start=1):
                    note(f"{block} tiers", tier, f"{key} path {p} T{t}")
    for key, enemy in data["Enemies"].items():
        note("Enemies", enemy, key)
    for key, difficulty in data["Difficulties"].items():
        note("Difficulties", difficulty, key)
    for i, row in enumerate(data["Mastery"], start=1):
        note("Mastery", row, f"level {i}")
    for key, perks in data["MasteryPerks"].items():
        for perk in perks:
            note("MasteryPerks", perk, f"{key} level {perk['level']}")
    for i, row in enumerate(data["Rounds"], start=1):
        note("Rounds", row, f"round {i}")
        fields.setdefault(("Rounds", "counts"), []).append(f"round {i}")
    for i, row in enumerate(data["DailyHaul"], start=1):
        note("DailyHaul", row, f"day {i}")
    for key, bounty in data["Bounties"].items():
        note("Bounties", bounty, key)
    for item in data["Cosmetics"]:
        note("Cosmetics", item, item["id"])
    for arena in data["Arenas"]:
        note("Arenas", arena, arena["key"])
    return fields


def tuning_formula_uses():
    """{Tuning field: how many workbook formulas read its cell}. A Tuning lever that
    only feeds sheet formulas (cost curves, round HP) shows up as a dead column; this
    tells those apart for whoever decides what to do with them."""
    from openpyxl import load_workbook
    wb = load_workbook(export_constants.DEFAULT_XLSX)
    ws = wb["Tuning"]
    rows = {}
    for r in range(5, 200):
        name, value = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if name and isinstance(value, (int, float)):
            rows[export_constants.ident(name)] = r
    formulas = []  # (sheet title, formula text)
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                text = getattr(cell.value, "text", cell.value)
                if isinstance(text, str) and text.startswith("="):
                    formulas.append((sheet.title, text))
    uses = {}
    for field, r in rows.items():
        other = re.compile(r"Tuning'?!\$?B\$?" + str(r) + r"(?!\d)")
        local = re.compile(r"(?<![A-Za-z$!])\$?B\$?" + str(r) + r"(?!\d)")
        uses[field] = sum(1 for title, text in formulas
                          if other.search(text) or (title == "Tuning" and local.search(text)))
    return uses


def check_dead(findings, data, notes):
    tuning_uses = tuning_formula_uses()
    config = os.path.join(ROOT, "src", "shared", "Config.luau")
    code = ""
    for path in glob.glob(os.path.join(ROOT, "src", "**", "*.luau"), recursive=True):
        if os.path.abspath(path) != config:
            with open(path, encoding="utf-8") as f:
                code += f.read() + "\n"
    for (block, field), where in sorted(config_fields(data).items()):
        used = re.search(r"\b" + re.escape(field) + r"\b", code)
        if used and (block, field) in DEAD_PENDING:
            findings.append(f"DEAD_PENDING lists {block}.{field}, which a src file now reads: remove the entry")
        if where and not used:
            uses = tuning_uses.get(field, 0) if block == "Tuning" else 0
            if (block, field) in DEAD_PENDING:
                notes.append(f"pending: {block}.{field} (read from {DEAD_PENDING[(block, field)]})")
            elif (block, field) in DEAD_ALLOWED:
                notes.append(f"allowed: {block}.{field} ({DEAD_ALLOWED[(block, field)]})")
            elif uses:
                notes.append(f"sheet-only: {block}.{field} (sheet formulas read it {uses}x)")
            else:
                shown = ", ".join(where[:3]) + (f" (+{len(where) - 3} more)" if len(where) > 3 else "")
                finding = f"{block}.{field} is set ({shown}) but no src file reads it"
                if block == "Tuning":
                    finding += "; no sheet formula reads it either"
                findings.append(finding)
    for (block, field) in DEAD_ALLOWED:
        if (block, field) not in config_fields(data):
            findings.append(f"DEAD_ALLOWED lists {block}.{field}, which Config no longer has")
    for (block, field) in DEAD_PENDING:
        if (block, field) not in config_fields(data):
            findings.append(f"DEAD_PENDING lists {block}.{field}, which Config no longer has")


# ---------------------------------------------------------------- stakes (phase 7)

def check_stakes(findings, notes, data, vision=None):
    """VISION.md "Amber economy": the 10-20 buy-in and the 72/23/5 Royale split equal the
    exporter's range and the Tuning levers (PLAN phase 7 T73)."""
    text = vision if vision is not None else read("VISION.md")
    tuning = data["Tuning"]
    m = re.search(r"(\d+)\s*[\u2013-]\s*(\d+) Core buy-in", text)
    if not m:
        findings.append("VISION.md: no 'N-M Core buy-in' line to check the buy-in against")
    else:
        low, high = int(m.group(1)), int(m.group(2))
        if (low, high) != export_constants.BUY_IN_RANGE:
            findings.append(f"VISION.md buy-in {low}-{high} != exporter range {export_constants.BUY_IN_RANGE}")
        buy_in = tuning.get("CompetitiveBuyIn")
        if buy_in is None or not (low <= buy_in <= high):
            findings.append(f"Tuning CompetitiveBuyIn ({buy_in}) is outside VISION.md's {low}-{high}")
        else:
            notes.append(f"buy-in {buy_in:g} within VISION.md {low}-{high}")
    m = re.search(r"1st gets \*\*(\d+)%\*\*, 2nd \*\*(\d+)%\*\*, 3rd \*\*(\d+)%\*\*", text)
    if not m:
        findings.append("VISION.md: no '1st gets **N%**, 2nd **N%**, 3rd **N%**' line to check the Royale split against")
    else:
        want = [int(g) for g in m.groups()]
        have = [tuning.get(k) for k in ("RoyalePot1st", "RoyalePot2nd", "RoyalePot3rd")]
        if want != have:
            findings.append(f"Tuning Royale pot {have} != VISION.md {want}")
        else:
            notes.append(f"Royale split {'/'.join(map(str, want))} = VISION.md")


# ---------------------------------------------------------------- 4. species

def check_species(findings, data):
    enemies = data["Enemies"]
    order = data["EnemyOrder"]
    by_display = {enemy["display"]: key for key, enemy in enemies.items()}
    rows = []
    in_table = False
    for line in read("VISION.md").splitlines():
        if line.startswith("| Tier | Species"):
            in_table = True
        elif in_table and line.startswith("|"):
            row = cells(line)
            if not set(row[0]) <= set("-"):
                rows.append((row[0], plain(row[1])))
        elif in_table:
            break
    if not rows:
        findings.append("VISION.md has no '| Tier | Species |' table")
        return
    listed = set()
    for tier, name in rows:
        key = by_display.get(name)
        if key is None:
            findings.append(f"VISION.md species '{name}' (tier {tier}) is no enemy display name in the sheet")
            continue
        listed.add(key)
        enemy = enemies[key]
        if tier.isdigit() and (int(tier) > len(order) or order[int(tier) - 1] != key):
            findings.append(f"VISION.md puts {name} at tier {tier}; the sheet has it as Enemies.{key} (#{order.index(key) + 1})")
        elif tier == "air" and not enemy.get("flying"):
            findings.append(f"VISION.md says {name} flies; Enemies.{key} isn't flying")
        elif tier == "boss" and not enemy.get("boss"):
            findings.append(f"VISION.md says {name} is a boss; Enemies.{key} isn't")
    for key, enemy in enemies.items():
        if key not in listed:
            findings.append(f"Enemies.{key} ({enemy['display']}) isn't in VISION.md's species table")


# ---------------------------------------------------------------- 5. mastery perks

def heroes_mastery_section():
    """The text of HEROES.md's "## Mastery" section."""
    match = re.search(r"^## Mastery[^\n]*\n(.*?)(?=^## |\Z)", read("HEROES.md"), re.S | re.M)
    return match.group(1) if match else ""


def taken_names(data):
    """{lower-cased name: where it is used} for every tier name on Tower Upgrades and
    Hero Upgrades and every hero's ability name: the names a perk may not reuse."""
    taken = {}
    for block in ("Towers", "Heroes"):
        for key, item in data[block].items():
            for path in item["paths"]:
                for t, tier in enumerate(path["tiers"], start=1):
                    name = str(tier.get("name") or "").strip()
                    if name:
                        taken.setdefault(name.lower(), f"{block}.{key} {path['id']} tier {t}")
    for key, hero in data["Heroes"].items():
        name = str(hero.get("ability") or "").strip()
        if name:
            taken.setdefault(name.lower(), f"Heroes.{key}'s ability")
    return taken


def check_perks(findings, notes, data, doc=None):
    perks = data["MasteryPerks"]
    doc = heroes_mastery_section() if doc is None else doc
    fields = [field for field, _ in export_constants.PERK_COLUMNS]
    taken = taken_names(data)
    undocumented, count = [], 0
    for key, rows in perks.items():
        for perk in rows:
            count += 1
            where = f"MasteryPerks.{key} level {perk['level']} ({perk['name']})"
            if perk["name"].lower() in taken:
                findings.append(f"{where}: the name is already {taken[perk['name'].lower()]}; "
                                "a perk needs a name of its own (DECISIONS #104)")
            allowed = [perk[f] for f in fields if isinstance(perk[f], (int, float)) and perk[f]]
            if isinstance(perk["markExtraPowerMult"], (int, float)):
                allowed.append(perk["markExtraPowerMult"] * 100)
            for number in re.findall(r"\d+(?:\.\d+)?", perk["text"]):
                if not any(abs(float(number) - value) < 1e-6 for value in allowed):
                    findings.append(f"{where}: Text says {number}, which is none of the row's effect cells")
            if "pop" in perk["text"].lower():
                findings.append(f"{where}: Text says \"pop\"; players read take-downs (DECISIONS #93)")
            if perk["name"] and perk["name"] not in doc:
                undocumented.append(perk["name"])
    for field, head in export_constants.PERK_COLUMNS:
        if count and not any(live(perk[field]) for rows in perks.values() for perk in rows):
            findings.append(f"Mastery Perks column '{head}' is blank on every row (dead column)")
    if undocumented:
        notes.append(f"not yet in HEROES.md \"Mastery\" ({len(undocumented)} of {count} perk names; PLAN round 3 T39 adds them): "
                     + ", ".join(undocumented))
    else:
        notes.append(f"{count} perk names found in HEROES.md \"Mastery\"")


# ---------------------------------------------------------------- 6. mastery columns

# Words a mastery column may never touch (PLAN T51, DECISIONS #135), matched on the
# header text and on the Config field name, both lowercased with spaces and
# punctuation removed.
MASTERY_FORBIDDEN = ("damage", "firerate", "rate", "reload", "recoil", "spread", "maxhp", "maxhealth", "health", "speed", "walk")


def check_mastery_columns(findings, notes, data):
    import openpyxl
    ws = openpyxl.load_workbook(export_constants.DEFAULT_XLSX, read_only=True)["Mastery"]
    heads = [str(c.value) for c in next(ws.iter_rows(min_row=4, max_row=4)) if c.value is not None]
    fields = sorted({f for row in data["Mastery"] for f in row})
    for where, names in (("header", heads), ("Config.Mastery field", fields)):
        for name in names:
            flat = re.sub(r"[^a-z]", "", name.lower())
            hit = [w for w in MASTERY_FORBIDDEN if w in flat]
            if hit:
                findings.append(f"Mastery {where} '{name}' touches {', '.join(hit)} (mastery is never combat power, DECISIONS #135)")
    notes.append(f"{len(heads)} Mastery headers and {len(fields)} fields checked: none touches damage, fire rate, reload, recoil, spread, max HP or move speed")


# ---------------------------------------------------------------- main

# Bounty events (PLAN round 2 T28): every event in Shared/Bounties EVENTS must be fired
# from server code, through Progression's one entry point (bountyEvent), the hook Shop
# and Hero are handed (onBounty), or Progression's own match-end call (bountyProgress);
# and nothing may fire an event that isn't in EVENTS.
BOUNTY_CALL = re.compile(r'\b(?:bountyEvent|onBounty|bountyProgress)\(\s*[\w.]+\s*,\s*"(\w+)"')


def bounty_call_sites(sources):
    """{event: [file, ...]} for every literal event name fired in `sources` ({file: text})."""
    sites = {}
    for name, text in sorted(sources.items()):
        for event in BOUNTY_CALL.findall(text):
            sites.setdefault(event, []).append(name)
    return sites


def check_bounty_events(findings, notes, sources=None):
    match = re.search(r"Bounties\.EVENTS = table\.freeze\(\{([^}]*)\}\)", read("src/shared/Bounties.luau"))
    if not match:
        findings.append("src/shared/Bounties.luau has no Bounties.EVENTS list")
        return
    events = re.findall(r'"(\w+)"', match.group(1))
    if sources is None:
        sources = {}
        for path in glob.glob(os.path.join(ROOT, "src", "server", "*.luau")):
            sources[os.path.basename(path)] = read(os.path.relpath(path, ROOT))
    sites = bounty_call_sites(sources)
    for event in events:
        if event not in sites:
            findings.append(f"bounty event '{event}' (Bounties.EVENTS) is never fired from src/server")
    for event, files in sites.items():
        if event not in events:
            findings.append(f"{files[0]} fires bounty event '{event}', which isn't in Bounties.EVENTS")
    if not findings:
        notes.append("bounty events fired: " + "; ".join(f"{e} ({', '.join(sorted(set(sites[e])))})" for e in events))


def main():
    strict = "--strict" in sys.argv[1:]
    with contextlib.redirect_stdout(io.StringIO()):
        data = export_constants.read_data(export_constants.DEFAULT_XLSX)

    compared = []
    checks = [
        ("names", lambda f, n: (check_names(f, "Towers", data["Towers"], "UPGRADES.md", "##", compared),
                             check_names(f, "Heroes", data["Heroes"], "HEROES.md", "###", compared))),
        ("unlocks", lambda f, n: check_unlocks(f, data)),
        ("dead columns", lambda f, n: check_dead(f, data, n)),
        ("species", lambda f, n: check_species(f, data)),
        ("bounty events", lambda f, n: check_bounty_events(f, n)),
        ("mastery perks", lambda f, n: check_perks(f, n, data)),
        ("mastery columns", lambda f, n: check_mastery_columns(f, n, data)),
        ("stakes", lambda f, n: check_stakes(f, n, data)),
    ]
    total = 0
    for name, run in checks:
        findings, notes = [], []
        run(findings, notes)
        total += len(findings)
        counted = f" ({len(compared)} tier names compared)" if name == "names" else ""
        print(f"{'ok  ' if not findings else 'FIND'} audit {name}: {len(findings)} finding(s){counted}")
        for note in notes:
            print(f"       ok {note}")
        for finding in findings:
            print(f"       - {finding}")
    print(f"audit: {total} finding(s)")
    sys.exit(1 if strict and total else 0)


if __name__ == "__main__":
    main()
